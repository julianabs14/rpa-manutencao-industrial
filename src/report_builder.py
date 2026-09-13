import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from src.config import PASTA_RELATORIOS

def gerar_relatorio_excel(resumo_df):
    os.makedirs(PASTA_RELATORIOS, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo de Manutencao"

    cabecalho_fill = PatternFill(start_color="7A2E3A", end_color="7A2E3A", fill_type="solid")
    cabecalho_font = Font(color="FFFFFF", bold=True)
    critico_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

    colunas = list(resumo_df.columns)
    ws.append(colunas)

    for celula in ws[1]:
        celula.fill = cabecalho_fill
        celula.font = cabecalho_font
        celula.alignment = Alignment(horizontal="center")

    for _, linha in resumo_df.iterrows():
        ws.append(list(linha))

    indice_coluna_critico = colunas.index("critico") + 1
    for numero_linha in range(2, ws.max_row + 1):
        celula_critico = ws.cell(row=numero_linha, column=indice_coluna_critico)
        if celula_critico.value is True:
            for coluna in range(1, len(colunas) + 1):
                ws.cell(row=numero_linha, column=coluna).fill = critico_fill

    for indice, coluna in enumerate(colunas, start=1):
        letra = get_column_letter(indice)
        maior_valor = max(
            [len(str(coluna))] + [len(str(v)) for v in resumo_df[coluna]]
        )
        ws.column_dimensions[letra].width = maior_valor + 4

    nome_arquivo = f"relatorio_manutencao_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    caminho_completo = os.path.join(PASTA_RELATORIOS, nome_arquivo)
    wb.save(caminho_completo)

    return caminho_completo
    
